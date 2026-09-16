"""Build a six-page IEEE-style paper draft from verified SecureLLMBench artifacts."""
from __future__ import annotations

import json
import math
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "SecureLLMBench_IEEE_Paper_Draft.docx"
ASSETS = ROOT / "artifacts" / "paper_assets"
TRACE = ROOT / ".tmp" / "paper_trace_real_local.json"
ANALYSIS = ROOT / ".tmp" / "paper_analysis_real_local.json"
BLUE = "#17365D"
OOXML_BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)


def set_columns(section, num=2, space=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space))
    if not sect_pr.xpath("./w:cols"):
        sect_pr.append(cols)


def font_for(size=30, bold=False):
    for name in (
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        "Times New Roman Bold.ttf" if bold else "Times New Roman.ttf",
        "DejaVuSerif-Bold.ttf" if bold else "DejaVuSerif.ttf",
    ):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline=BLUE):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 20, y1 + 14), title, fill=BLUE, font=font_for(25, True))
    y = y1 + 54
    for line in lines:
        draw.text((x1 + 20, y), line, fill="#222222", font=font_for(20))
        y += 28


def arrow(draw, a, b):
    draw.line([a, b], fill=BLUE, width=5)
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    wing = 15
    for delta in (2.6, -2.6):
        p = (b[0] - wing * math.cos(angle + delta), b[1] - wing * math.sin(angle + delta))
        draw.line([b, p], fill=BLUE, width=5)


def create_architecture():
    im = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 24), "SecureLLMBench traceable evaluation architecture", fill=BLUE, font=font_for(33, True))
    boxes = [
        ((60, 150, 335, 350), "Versioned inputs", ["Dataset hash", "Prompt case", "Generation config"], "#EAF3F8"),
        ((430, 150, 710, 350), "Execution", ["Ollama / mock", "Target response", "Latency + metadata"], "#E8F1FC"),
        ((805, 80, 1105, 275), "Measurement", ["Layer 1 evidence", "Layer 2 judge", "Behavioral state"], "#F2F7FB"),
        ((805, 410, 1105, 635), "Metrics", ["BSDA components", "RC trajectory", "SAEA diagnostics"], "#EAF3F8"),
        ((1200, 220, 1450, 470), "Governed outputs", ["Records + API", "Dashboard", "Explicit null states"], "#E8F1FC"),
    ]
    for b in boxes:
        rounded_box(d, *b)
    arrow(d, (335, 250), (430, 250))
    arrow(d, (710, 220), (805, 180))
    arrow(d, (710, 290), (805, 500))
    arrow(d, (1105, 180), (1200, 300))
    arrow(d, (1105, 520), (1200, 390))
    d.text((70, 690), "Every edge carries provenance; unavailable measurements remain unavailable instead of being converted to zero.", fill="#555555", font=font_for(22))
    im.save(ASSETS / "architecture.png")


def create_latency_chart(latencies):
    labels = ["Baseline", "Isolated", "Sequence 1", "Sequence 2", "Recovery"]
    im = Image.new("RGB", (1500, 800), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 28), "Local Ollama engineering trace: response latency", fill=BLUE, font=font_for(34, True))
    left, top, right, bottom = 125, 135, 1440, 650
    d.line((left, top, left, bottom), fill="#333333", width=3)
    d.line((left, bottom, right, bottom), fill="#333333", width=3)
    maximum = 1700
    for value in range(0, 1800, 400):
        y = bottom - (value / maximum) * (bottom-top)
        d.line((left, y, right, y), fill="#DDDDDD", width=2)
        d.text((25, y-11), str(value), fill="#555555", font=font_for(19))
    bar_w, gap = 150, 105
    for i, value in enumerate(latencies):
        x = 190 + i * (bar_w + gap)
        y = bottom - (value / maximum) * (bottom-top)
        d.rounded_rectangle((x, y, x+bar_w, bottom), radius=8, fill="#3F78A8")
        d.text((x+9, y-35), f"{value:.0f}", fill=BLUE, font=font_for(20, True))
        d.text((x, bottom+18), labels[i], fill="#333333", font=font_for(18))
    d.text((640, 720), "Latency (ms); one local engineering trace, n = 5 stages", fill="#555555", font=font_for(21))
    im.save(ASSETS / "latency.png")


def create_status_matrix():
    im = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 28), "Current evidence status by measurement family", fill=BLUE, font=font_for(34, True))
    rows = [
        ("Layer 1", "Evidence available", "Not attack success"),
        ("Layer 2", "Structured measurement", "Judge validation required"),
        ("BSDA", "Formula/runtime ready", "Empirical validation required"),
        ("Recovery Capability", "Trajectory stored", "Calibration artifact required"),
        ("SAEA", "Diagnostics ready", "Matched trajectories required"),
        ("DRAA / PRI", "Evidence profiles", "Scalar intentionally unavailable"),
    ]
    x = [55, 430, 845, 1440]
    y0, row_h = 115, 88
    headers = ["Family", "Current state", "Claim boundary"]
    for i in range(3):
        d.rectangle((x[i], y0, x[i+1], y0+row_h), fill=BLUE)
        d.text((x[i]+18, y0+24), headers[i], fill="white", font=font_for(22, True))
    for r, row in enumerate(rows):
        y = y0 + (r+1)*row_h
        fill = "#F2F7FB" if r % 2 else "#FFFFFF"
        for i in range(3):
            d.rectangle((x[i], y, x[i+1], y+row_h), fill=fill, outline="#D0D7DF", width=2)
            d.text((x[i]+18, y+24), row[i], fill="#222222", font=font_for(20))
    d.text((60, 790), "The dashboard reports the state and provenance of each value; it does not turn these boundaries into an aggregate score.", fill="#555555", font=font_for(19))
    im.save(ASSETS / "status_matrix.png")


def set_run_font(run, size=9.2, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_p(doc, text="", size=9.2, align=None, before=0, after=3, first_indent=True, bold_prefix=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if first_indent:
        pf.first_line_indent = Inches(0.14)
    if align:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size)
    else:
        r = p.add_run(text)
        set_run_font(r, size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 10.2 if level == 1 else 9.5, bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, 8.2, italic=True)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, OOXML_BLUE)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 7.6, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margin(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, 7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("SecureLLMBench: Traceable LLM Security Evaluation")
    set_run_font(r, 8.0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("IEEE-style paper draft | repository evidence through 11 September 2026")
    set_run_font(r, 7.5)


def load_trace():
    trace = json.loads(TRACE.read_text(encoding="utf-8"))
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    responses = trace["responses"]
    latency = [responses["baseline"]["latency_ms"], responses["isolated_attack"]["latency_ms"]]
    latency.extend(r["latency_ms"] for r in responses["sequential"])
    latency.append(responses["recovery"][0]["latency_ms"])
    return trace, analysis, latency


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    trace, analysis, latencies = load_trace()
    create_architecture()
    create_latency_chart(latencies)
    create_status_matrix()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    add_header_footer(section)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")

    # Page 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("SecureLLMBench: A Traceable Framework for Safety and Security Evaluation of Local Large Language Models")
    set_run_font(r, 14.5, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Anonymous IEEE-Style Research Draft")
    set_run_font(r, 9.5, italic=True)
    add_heading(doc, "Abstract", 1)
    add_p(doc, "SecureLLMBench is a research-oriented software framework for conducting traceable evaluations of local large language models (LLMs) under benign, isolated-attack, sequential-attack, and recovery conditions. The framework persists versioned inputs, generation settings, target-model metadata, detector evidence, structured judge measurements, metric artifacts, and explicit failure states. This paper documents the implemented architecture and a local engineering trace using qwen3.5:2b as target and gemma3:4b as judge. The trace confirms end-to-end execution and provenance capture, but it is not an empirical security result: target responses were empty at the configured generation limit, some judge dimensions failed validation, and attack-success labels were absent. The contribution is therefore a reproducible measurement and governance substrate, not a claim that a model is secure or that any proposed metric is validated.")
    add_p(doc, "Index Terms—LLM security evaluation, reproducibility, local inference, benchmark provenance, behavioral measurement, scientific safeguards.", size=8.7, first_indent=False, bold_prefix="Index Terms—")
    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body_section)
    add_heading(doc, "I. INTRODUCTION", 1)
    add_p(doc, "Security evaluation for LLMs requires more than a prompt and a score. A defensible result must identify the benchmark population, model version, generation settings, prompt sequence, measurements, evaluator behavior, and all unavailable prerequisites. SecureLLMBench addresses this engineering requirement by treating an evaluation as an auditable trace rather than as an opaque number. It connects dataset ingestion, local inference, Layer 1 detector evidence, Layer 2 structured judging, behavioral metrics, persistent records, a FastAPI backend, and a research dashboard.")
    add_p(doc, "The central design choice is semantic restraint. Layer 1 outputs are detector-native evidence rather than attack-success truth. Layer 2 outputs are structured measurements rather than calibrated probabilities. Missing, failed, uncalibrated, undefined, and not-applicable values remain distinct values in storage and in the dashboard. This prevents a user interface or downstream report from implying that a missing measurement is evidence of zero risk.")
    add_heading(doc, "A. Contributions", 2)
    add_p(doc, "First, the framework implements an end-to-end local evaluation trace with database persistence, read-only API endpoints, and dashboard exploration. Second, it implements four-component BSDA computation, Recovery Capability (RC), and Sequential Attack Evaluation Algorithm (SAEA) runtimes with explicit applicability contracts. Third, it preserves DRAA and PRI as uncalibrated evidence profiles, deliberately withholding scalar risk and model-ranking values. Fourth, it provides reproducibility artifacts, safety tests, and explicit research blockers.")
    add_heading(doc, "B. Scope of This Paper", 2)
    add_p(doc, "This paper is an implementation and trace report. It does not present a broad benchmark comparison, validated attack success rate (ASR), calibration study, or model ranking. The included real local trace contains one repository example case and is used only to demonstrate system behavior and to show how scientific uncertainty is carried forward.")
    page_break(doc)

    # Page 2
    add_heading(doc, "II. SYSTEM ARCHITECTURE", 1)
    add_p(doc, "Figure 1 summarizes the evaluation flow. A case begins with a versioned dataset record and a captured generation configuration. A provider produces a response together with timing and model identity. Layer 1 detectors provide rule, keyword, and pattern evidence. Layer 2 produces dimension-specific structured judgments. The framework maps compatible judgments into behavioral states, computes metric artifacts, and persists both values and their provenance for dashboard and API consumption.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "architecture.png"), width=Inches(6.75))
    add_caption(doc, "Fig. 1. Traceable evaluation flow. Data and state move forward with configuration, provenance, and explicit availability semantics.")
    add_heading(doc, "A. Data and Execution Contracts", 2)
    add_p(doc, "The data layer ingests JSON, JSONL, and CSV sources and computes a dataset hash before analysis. The execution layer supports deterministic mock inference for software tests and local Ollama inference for already-installed models. No experiment command downloads a model. Each provider result stores model/provider metadata, latency, finish reason, generation settings, and raw response metadata so that a later reviewer can distinguish a successful response from a length-limited or failed output.")
    add_table(doc, ["Artifact", "Captured fields", "Reason"], [
        ["Dataset version", "Path, hash, case id, taxonomy", "Recreates the evaluated population"],
        ["Model execution", "Model digest, seed, temperature, token limit, latency", "Supports configuration-level comparison"],
        ["Layer 1", "Detector score, evidence, rule metadata", "Preserves detector-native semantics"],
        ["Layer 2", "Dimension, rubric/prompt versions, status, confidence", "Makes evaluator behavior auditable"],
        ["Metric record", "Formula inputs, applicability, provenance", "Prevents unsupported aggregation"],
    ], [1.12, 2.75, 2.72])
    add_caption(doc, "TABLE I. Core trace artifacts and their audit purpose.")
    add_heading(doc, "B. Dashboard and API", 2)
    add_p(doc, "The backend exposes safe read-only endpoints for dashboard summary, datasets, models, benchmark runs, evaluations, Layer 1 and Layer 2 results, scientific records, model-score provenance, runtime models, and review records. The dashboard renders these items in a light research UI and differentiates empty data from unavailable upstream measurements. Local execution is exposed as an opt-in capability; a server restart does not manufacture benchmark jobs or synthetic results.")
    page_break(doc)

    # Page 3
    add_heading(doc, "III. MEASUREMENT DESIGN AND CALCULATIONS", 1)
    add_p(doc, "SecureLLMBench uses formulas only within their stated data contracts. The formulas below define computational behavior; they are not claimed as validated scientific estimators. Every metric requires compatible input scales, declared versions, and an explicit applicability outcome.")
    add_heading(doc, "A. Dataset Quality Index", 2)
    add_p(doc, "The exploratory Dataset Quality Index (DQI) is a weighted mean of duplicate quality, prompt coverage, category entropy, novelty, difficulty entropy, and category balance. With component values c_k and non-negative weights w_k, the implemented calculation is DQI = sum_k w_k c_k / sum_k w_k. Novelty is zero when no embeddings are supplied; this is an implementation convention and not evidence that the dataset lacks novelty. DQI must therefore be reported as exploratory only.")
    add_heading(doc, "B. Recovery Capability", 2)
    add_p(doc, "For an active behavioral vector with m dimensions, the primary RC distance is d(X,Y) = sqrt[(1/m) sum_j (X_j - Y_j)^2]. Given baseline mean Bbar, attack state A, and recovery states R_t, the attack displacement is Delta_A = d(A,Bbar), and each recovery displacement is Delta_t = d(R_t,Bbar). The primary raw score is RC_AUC_raw = 1 - mean_t(Delta_t)/Delta_A. Terminal recovery is 1 - Delta_n/Delta_A. Both are null unless safety and helpfulness are available, context is retained, and a compatible external calibration artifact supplies a positive Delta_A applicability floor.")
    add_heading(doc, "C. Sequential Attack Evaluation", 2)
    add_p(doc, "For ordered states s_i, SAEA retains per-step deviation Delta_i = d(s_i,Bbar), observed cumulative deviation CV_obs = (1/L) sum_i Delta_i, and the candidate Bliss null E_Bliss = 1 - product_i(1 - Delta_i) under normalized Euclidean Delta_max = 1. A sequence-context effect DeltaV_i = Delta_i - V_iso(A_i) is allowed only when the isolated comparison is commensurate and matched. The framework never zero-fills a missing dimension or treats an undefined interaction as no interaction.")
    add_table(doc, ["Family", "Permitted calculation", "Current output state"], [
        ["BSDA", "Four separate components: semantic, safety, instruction, structural", "No composite score"],
        ["RC", "Normalized Euclidean recovery trajectory", "Null if uncalibrated/incompatible"],
        ["SAEA", "Per-step, CV_obs, candidate null diagnostics", "Requires matched trajectories"],
        ["DRAA", "Evidence transport and missingness provenance", "risk_score = null"],
        ["PRI", "Model/configuration evidence profile", "scalar_pri = null"],
    ], [0.72, 3.16, 2.71])
    add_caption(doc, "TABLE II. Metric families and enforced output boundaries.")
    add_heading(doc, "D. Outcome and Human Validation", 2)
    add_p(doc, "ASR is calculated only from explicit classified outcome records: ASR = successes / classified denominator. Layer 1 or Layer 2 values cannot silently become success labels. The repository includes blinded sample export, agreement rate, Cohen’s kappa, Fleiss kappa, and confusion-matrix utilities for a future independent rater protocol. These utilities are workflow support, not evidence that human validation has occurred.")
    page_break(doc)

    # Page 4
    add_heading(doc, "IV. LOCAL ENGINEERING TRACE", 1)
    add_p(doc, "A real local trace was executed with qwen3.5:2b as the target model and gemma3:4b as the local judge. The trace used the repository example dataset hash b395724a958530e590e2800a9a3c351fb702c6bb0b36327cb81070d0b8375c1c, temperature 0.0, seed 2026, and a 128-token generation limit. It evaluated a baseline task, an isolated injection prompt, a two-step sequence, and a recovery instruction. The purpose of the trace was to verify runtime integration, not to estimate security performance.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "latency.png"), width=Inches(6.75))
    add_caption(doc, "Fig. 2. Stage-level latency from the local trace. Mean across five stages: 1448.0 ms; sample standard deviation: 85.3 ms.")
    add_heading(doc, "A. Observed Runtime Values", 2)
    add_table(doc, ["Stage", "Latency (ms)", "Finish reason", "Layer 1 aggregate"], [
        ["Baseline", f"{latencies[0]:.1f}", "length", "0.0"],
        ["Isolated attack", f"{latencies[1]:.1f}", "length", "0.0"],
        ["Sequence 1", f"{latencies[2]:.1f}", "length", "0.0"],
        ["Sequence 2", f"{latencies[3]:.1f}", "length", "0.0"],
        ["Recovery", f"{latencies[4]:.1f}", "length", "0.0"],
    ], [1.45, 1.25, 1.35, 1.35])
    add_caption(doc, "TABLE III. Local engineering trace outputs. A detector aggregate of 0.0 is evidence from configured detectors, not attack failure or model safety.")
    add_heading(doc, "B. Interpretation", 2)
    add_p(doc, "All five target stages returned an empty text field with finish reason length. The judge produced structured output for some instruction-following dimensions but failed validation for safety-stance dimensions because the returned dimension did not match the request. The pipeline therefore recorded failed behavioral states, null RC scores with reason missing_required_behavioral_state_dimension, undefined SAEA values with reason missing_required_dimension, null DRAA risk, null PRI scalar, and ASR = null with a zero classified denominator. These are correct system outcomes under the framework’s contracts.")
    add_p(doc, "The trace confirms that the evaluator does not convert runtime incompleteness into an apparently favorable score. It does not support a claim about qwen3.5:2b security, gemma3:4b judge accuracy, benchmark-level effectiveness, or metric validity.")
    page_break(doc)

    # Page 5
    add_heading(doc, "V. SCIENTIFIC-INTEGRITY CONTROLS", 1)
    add_p(doc, "A research dashboard becomes misleading when it collapses operational status into a score. SecureLLMBench uses a typed availability model so that the user can distinguish not evaluated, not applicable, failed, undefined, unavailable upstream, insufficient data, uncalibrated, and computed values. This model is present in persisted records and frontend components, and it is maintained through metric composition.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "status_matrix.png"), width=Inches(6.75))
    add_caption(doc, "Fig. 3. Evidence status reported by the current implementation. These states delimit what can be claimed from the available artifacts.")
    add_heading(doc, "A. Prohibited Shortcuts", 2)
    add_p(doc, "The implementation prohibits scalar BSDA composites, DRAA risk scores without a calibration artifact, scalar PRI values, unlabelled ASR, and risk probabilities inferred from bounded judge scores. It also prevents RC from treating blanket refusal as recovery by requiring both safety and helpfulness dimensions. The primary RC result requires retained context, a compatible calibrated applicability floor, and a measurable attack displacement. SAEA keeps interaction and sequence diagnostics separate from a universal security score.")
    add_heading(doc, "B. Claim Matrix", 2)
    add_table(doc, ["Potential claim", "Evidence now", "Decision"], [
        ["Software traceability", "Versioned records, tests, API, dashboard", "Supported engineering claim"],
        ["Local execution works", "Ollama trace and smoke checks", "Supported engineering claim"],
        ["Model is secure", "One incomplete trace; no classified outcomes", "Not supported"],
        ["Metric is valid", "Formula/runtime tests only", "Not supported"],
        ["Model ranking", "No comparable validated benchmark population", "Not supported"],
        ["Risk probability", "No target labels or calibrated model", "Not supported"],
    ], [1.55, 2.88, 1.35])
    add_caption(doc, "TABLE IV. What the present artifacts do and do not substantiate.")
    add_heading(doc, "C. Threats to Validity", 2)
    add_p(doc, "The current dataset is a three-record repository engineering example, not a validated corpus. The local trace contains one case. Target outputs were empty, and the judge encountered structured-output validation failures. No independent outcome labels, rater adjudication, calibration set, or preregistered inferential plan exists. These limitations are persisted as blockers rather than hidden by aggregation or imputation.")
    page_break(doc)

    # Page 6
    add_heading(doc, "VI. REPRODUCIBILITY AND NEXT STUDY", 1)
    add_p(doc, "The repository contains a fixture dry-run, a local-Ollama runner, trace analysis scripts, unit tests, API smoke checks, and frontend typecheck, lint, test, and production-build commands. The verified software suite contains 129 passing Python tests and 19 passing frontend tests, together with successful typecheck, lint, build, API smoke, and local server checks. These checks establish implementation behavior; they do not validate a scientific hypothesis.")
    add_heading(doc, "A. Reproduction Procedure", 2)
    add_table(doc, ["Step", "Command / artifact", "Expected evidence"], [
        ["1", "python -m pytest", "Implementation tests"],
        ["2", "scripts/api_smoke.py", "Health and API contracts"],
        ["3", "run_experiment.py --fixture-dry-run", "Engineering-only trace JSON"],
        ["4", "run_experiment.py --target-model ... --judge-model ...", "Local trace for installed models"],
        ["5", "analyze_results.py --input trace.json", "Descriptive analysis plus blockers"],
        ["6", "frontend typecheck, lint, test, build", "Dashboard integrity"],
    ], [0.5, 3.2, 2.08])
    add_caption(doc, "TABLE V. Reproduction sequence. Each command records an engineering artifact rather than a scientific conclusion.")
    add_heading(doc, "B. Required Work Before a Research Claim", 2)
    add_p(doc, "A valid benchmark study needs a versioned threat model, a representative dataset with provenance, locked target and judge models, an approved outcome taxonomy and ASR denominator policy, independent labels or blinded human raters, a compatible RC calibration artifact, matched isolated and sequential trajectories, and a preregistered statistical plan. Any learned ML or DRAA scalar must be trained and validated against independently defined labels with grouped splits that prevent family, session, and near-duplicate leakage.")
    add_heading(doc, "VII. CONCLUSION", 1)
    add_p(doc, "SecureLLMBench provides a working foundation for local LLM safety and security evaluation: it runs models, captures a complete trace, computes permitted metric artifacts, preserves provenance, exposes safe read APIs, and renders research data in a dashboard. Its most important current output is disciplined uncertainty. The project is ready for a properly designed empirical study, but the present repository and local trace do not justify claims about model robustness, safety probabilities, metric validity, or rankings.")
    add_heading(doc, "REFERENCES", 1)
    refs = [
        "[1] SecureLLMBench, Repository implementation and reproducibility documentation, 2026.",
        "[2] SecureLLMBench, Recovery Capability Reconciled Implementation Specification, 2026.",
        "[3] SecureLLMBench, Sequential Attack Evaluation Algorithm Implementation Specification, 2026.",
        "[4] SecureLLMBench, DRAA and PRI Methodology-to-Implementation Reconciliations, 2026.",
        "[5] FastAPI, SQLAlchemy, React, and Ollama local runtime components, accessed in the implementation environment.",
    ]
    for ref in refs:
        add_p(doc, ref, size=8.2, first_indent=False, after=1)

    doc.core_properties.title = "SecureLLMBench Traceable LLM Security Evaluation"
    doc.core_properties.author = "SecureLLMBench"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
